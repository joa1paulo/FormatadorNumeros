"""Etapa 2: leitura em FLUXO de CSV e XLSX, formatação e remoção de duplicados.

O arquivo nunca é carregado inteiro na memória: as linhas são lidas e processadas uma
a uma. Só o resultado (números únicos) fica guardado, e de forma compacta.
"""
import codecs
import csv
import io
from dataclasses import dataclass, field
from itertools import chain, islice
from typing import NamedTuple

from formatador import formatar_numero

# palavra encontrada no cabeçalho -> prioridade (menor = preferida)
PALAVRAS_NUMERO = [
    ("whatsapp", 0), ("celular", 1),
    ("telefone", 2), ("phone", 2), ("fone", 2),
    ("numero", 3), ("número", 3),
]
LINHAS_AMOSTRA = 200
TAMANHO_BLOCO = 1 << 20  # 1 MB


class LimiteDeLinhas(ValueError):
    """O arquivo tem mais linhas do que o servidor aceita processar."""


class Descartado(NamedTuple):
    nome: str
    original: str
    motivo: str


@dataclass
class Resumo:
    total: int = 0
    duplicados: int = 0
    ok: dict = field(default_factory=dict)         # número E.164 -> nome (na ordem de chegada)
    fixos: list = field(default_factory=list)      # [Descartado]
    invalidos: list = field(default_factory=list)  # [Descartado]

    def como_lista(self) -> list:
        """[{nome, numero}] (útil em testes; evite em arquivos grandes)."""
        return [{"nome": nome, "numero": numero} for numero, nome in self.ok.items()]


@dataclass
class _Leitura:
    linhas: object          # iterador de linhas (listas de str), sem linhas vazias
    formato: str            # "csv" ou "xlsx"
    detalhe: str            # csv: o separador; xlsx: o nome da aba lida
    estimar_total: object   # () -> int | None  (conta/estima as linhas sem guardá-las)
    fechar: object          # () -> None


@dataclass
class Analise:
    formato: str
    detalhe: str
    tem_cabecalho: bool
    cabecalho: list
    idx_num: int           # coluna lida como número
    idx_nome: int | None   # coluna lida como nome (None se não houver)
    ncols: int
    amostra: list          # primeiras linhas de dados (já lidas)
    esgotado: bool         # True se o arquivo inteiro cabe na amostra
    leitura: _Leitura

    def dados(self):
        """Todas as linhas de dados, em fluxo: a amostra e depois o resto do arquivo."""
        return chain(self.amostra, self.leitura.linhas)


def detectar_delimitador(texto: str) -> str:
    """Descobre o separador pela 1ª linha: vírgula, ponto e vírgula ou tabulação.
    Ignora separadores dentro de aspas. Sem nenhum separador, assume vírgula."""
    linha = next((l for l in texto.splitlines() if l.strip()), "")
    contagem = {",": 0, ";": 0, "\t": 0}
    entre_aspas = False
    for ch in linha:
        if ch == '"':
            entre_aspas = not entre_aspas
        elif not entre_aspas and ch in contagem:
            contagem[ch] += 1
    melhor = max(contagem, key=contagem.get)
    return melhor if contagem[melhor] > 0 else ","


def _prioridade(cabecalho: str) -> int | None:
    c = cabecalho.strip().lower()
    if c in ("tel", "tel."):
        return 2
    for palavra, prioridade in PALAVRAS_NUMERO:
        if palavra in c:
            return prioridade
    return None


# ---------------------------------------------------------------- leitura em fluxo
def _abrir_binario(conteudo):
    if isinstance(conteudo, str):
        return io.BytesIO(conteudo.encode("utf-8"))
    if isinstance(conteudo, (bytes, bytearray)):
        return io.BytesIO(conteudo)
    conteudo.seek(0)  # já é um arquivo (ex.: o upload do FastAPI, que fica em disco)
    return conteudo


def _detectar_codificacao(f) -> str:
    """Passa pelo arquivo em blocos (sem guardá-lo) para saber se é UTF-8, UTF-16 ou cp1252."""
    inicio = f.read(4)
    f.seek(0)
    if inicio[:2] in (b"\xff\xfe", b"\xfe\xff"):  # "Texto Unicode" do Excel
        return "utf-16"
    decodificador = codecs.getincrementaldecoder("utf-8")()
    try:
        while bloco := f.read(TAMANHO_BLOCO):
            decodificador.decode(bloco)
        decodificador.decode(b"", final=True)
        codificacao = "utf-8-sig"
    except UnicodeDecodeError:  # CSV salvo pelo Excel em português costuma ser cp1252
        codificacao = "cp1252"
    f.seek(0)
    return codificacao


def _linhas_de_texto(f, codificacao: str):
    """Entrega o arquivo linha a linha (com o \\n), decodificando em blocos de 1 MB."""
    decodificador = codecs.getincrementaldecoder(codificacao)(errors="replace")
    resto = ""
    while True:
        bloco = f.read(TAMANHO_BLOCO)
        resto += decodificador.decode(bloco, final=not bloco)
        partes = resto.split("\n")
        resto = partes.pop()
        for parte in partes:
            yield parte + "\n"
        if not bloco:
            break
    if resto:
        yield resto


def _texto_celula(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))  # 5521994833600.0 -> 5521994833600
    return str(v)


def _leitura_csv(f) -> _Leitura:
    codificacao = _detectar_codificacao(f)
    cabeca = f.read(65536)
    f.seek(0)
    delimitador = detectar_delimitador(cabeca.decode(codificacao, errors="ignore"))

    def linhas():
        try:
            for linha in csv.reader(_linhas_de_texto(f, codificacao), delimiter=delimitador):
                if linha and any(c.strip() for c in linha):
                    yield linha
        except csv.Error as erro:
            raise ValueError(f"Não foi possível ler o CSV: {erro}") from erro

    def estimar_total():
        f.seek(0)
        n = 0
        while bloco := f.read(TAMANHO_BLOCO):
            n += bloco.count(b"\n")
        return n

    return _Leitura(linhas(), "csv", delimitador, estimar_total, lambda: None)


def _leitura_xlsx(f) -> _Leitura:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError("Leitura de .xlsx indisponível neste servidor (falta instalar o openpyxl).")
    try:
        wb = load_workbook(f, read_only=True, data_only=True)  # read_only: lê a planilha em fluxo
        aba = wb.worksheets[0]  # primeira aba
    except Exception:
        raise ValueError("Não foi possível abrir o .xlsx (arquivo corrompido ou protegido por senha).")

    def linhas():
        try:
            for linha in aba.iter_rows(values_only=True):
                texto = [_texto_celula(c) for c in linha]
                if any(c.strip() for c in texto):
                    yield texto
        except Exception as erro:
            raise ValueError("Não foi possível ler a planilha (arquivo corrompido).") from erro
        finally:
            wb.close()

    return _Leitura(linhas(), "xlsx", aba.title, lambda: aba.max_row, wb.close)


def _abrir_leitura(f) -> _Leitura:
    """Aceita CSV (, ; tab) e XLSX; recusa o resto com uma mensagem clara."""
    inicio = f.read(8)
    f.seek(0)
    if inicio.startswith(b"PK\x03\x04"):  # .xlsx é um zip
        return _leitura_xlsx(f)
    if inicio.startswith(b"\xd0\xcf\x11\xe0"):  # .xls antigo ou arquivo protegido por senha
        raise ValueError("Arquivo .xls antigo ou protegido por senha. Salve como .xlsx ou .csv e envie de novo.")
    cabeca = f.read(4096)
    f.seek(0)
    if b"\x00" in cabeca and cabeca[:2] not in (b"\xff\xfe", b"\xfe\xff"):
        raise ValueError("Formato de arquivo não suportado. Envie um arquivo .csv ou .xlsx.")
    return _leitura_csv(f)


# ---------------------------------------------------------------- análise e processamento
def analisar(conteudo) -> Analise | None:
    """Lê só o começo do arquivo para descobrir cabeçalho e colunas; o resto fica em fluxo."""
    leitura = _abrir_leitura(_abrir_binario(conteudo))
    amostra_bruta = list(islice(leitura.linhas, LINHAS_AMOSTRA + 1))
    if not amostra_bruta:
        leitura.fechar()
        return None

    primeira = amostra_bruta[0]
    tem_palavra = any(_prioridade(c) is not None for c in primeira)
    parece_dado = any(formatar_numero(c).status == "ok" for c in primeira)
    tem_cabecalho = tem_palavra or not parece_dado  # lista só com números não tem cabeçalho

    cabecalho = primeira if tem_cabecalho else []
    amostra = amostra_bruta[1:] if tem_cabecalho else amostra_bruta
    ncols = max(len(l) for l in amostra_bruta)

    def validos(i: int) -> int:  # quantos números válidos a coluna tem na amostra
        return sum(1 for l in amostra[:LINHAS_AMOSTRA] if i < len(l) and formatar_numero(l[i]).status == "ok")

    melhor_por_conteudo = max(range(ncols), key=lambda i: (validos(i), -i))
    candidatas = [(i, _prioridade(c)) for i, c in enumerate(cabecalho) if _prioridade(c) is not None]
    if candidatas:
        # cabeçalho manda (WhatsApp > Celular > Telefone > Número); empate: mais válidos, depois a da esquerda
        idx_num = min(candidatas, key=lambda t: (t[1], -validos(t[0]), t[0]))[0]
        # coluna escolhida sem nenhum número válido, mas outra tem: usa a que tem
        if validos(idx_num) == 0 and validos(melhor_por_conteudo) > 0:
            idx_num = melhor_por_conteudo
    else:
        idx_num = melhor_por_conteudo  # cabeçalho não ajuda (ex.: "ID,Contato"): decide pelo conteúdo

    idx_nome = None
    if tem_cabecalho:
        baixo = [c.strip().lower() for c in cabecalho]
        idx_nome = next((i for i, c in enumerate(baixo) if c.startswith("nome")), None)
        if idx_nome is None:
            idx_nome = next((i for i, c in enumerate(baixo) if "nome" in c and "sobrenome" not in c), None)
        if idx_nome == idx_num:
            idx_nome = None

    return Analise(leitura.formato, leitura.detalhe, tem_cabecalho, cabecalho, idx_num, idx_nome,
                   ncols, amostra, len(amostra_bruta) <= LINHAS_AMOSTRA, leitura)


def processar_csv(conteudo, max_linhas: int | None = None) -> Resumo:
    """conteudo: bytes, str ou arquivo aberto (CSV ou XLSX). Lê em fluxo, linha a linha."""
    resumo = Resumo()
    a = analisar(conteudo)
    if a is None:
        return resumo

    ok = resumo.ok
    try:
        for linha in a.dados():
            resumo.total += 1
            if max_linhas is not None and resumo.total > max_linhas:
                raise LimiteDeLinhas(f"O arquivo passa de {max_linhas} linhas.")
            bruto = linha[a.idx_num] if a.idx_num < len(linha) else ""
            nome = linha[a.idx_nome].strip() if a.idx_nome is not None and a.idx_nome < len(linha) else ""

            r = formatar_numero(bruto)
            if r.status == "ok":
                atual = ok.get(r.numero)
                if atual is None:
                    ok[r.numero] = nome
                else:
                    resumo.duplicados += 1
                    if not atual and nome:  # aproveita o nome da linha repetida
                        ok[r.numero] = nome
            elif r.status == "fixo":
                resumo.fixos.append(Descartado(nome, bruto, r.motivo))
            else:
                resumo.invalidos.append(Descartado(nome, bruto, r.motivo))
    finally:
        a.leitura.fechar()
    return resumo


def previa_csv(conteudo, n: int = 8) -> dict:
    """Dados para o preview da tela: mostra exatamente como o sistema vai ler o arquivo.
    Lê só o começo do arquivo; o total de linhas de arquivos grandes é uma estimativa."""
    a = analisar(conteudo)
    if a is None:
        return {"vazio": True}
    try:
        if a.esgotado:
            total, aproximado = len(a.amostra), False
        else:
            estimado = a.leitura.estimar_total()
            total = None if estimado is None else max(estimado - (1 if a.tem_cabecalho else 0), len(a.amostra))
            aproximado = True
    finally:
        a.leitura.fechar()

    ncols = max(a.ncols, len(a.cabecalho))
    cabecalho = [*a.cabecalho, *[f"Coluna {i + 1}" for i in range(len(a.cabecalho), ncols)]]
    linhas = [[*l, *[""] * (ncols - len(l))] for l in a.amostra[:n]]
    return {
        "vazio": False,
        "formato": a.formato,
        "delimitador": a.detalhe if a.formato == "csv" else "",
        "aba": a.detalhe if a.formato == "xlsx" else "",
        "tem_cabecalho": a.tem_cabecalho,
        "cabecalho": cabecalho,
        "linhas": linhas,
        "idx_numero": a.idx_num,
        "total_linhas": total,
        "aproximado": aproximado,
    }


def _testes():
    contatados = (
        "Nome do Contato,WhatsApp do Contato,Etapa da Jornada\n"
        "Lucia,5521994833600,Fez Contato\n"
        "Daniele ,5521978785757,Fez Contato\n"
        ",5521970121596,Fez Contato\n"
        "Lucia repetida,5521994833600,Fez Contato\n"
    )
    r = processar_csv(contatados)
    lista = r.como_lista()
    assert (r.total, len(r.ok), r.duplicados) == (4, 3, 1)
    assert lista[1]["nome"] == "Daniele"          # espaço removido
    assert lista[2]["nome"] == ""                 # nome vazio tolerado
    assert lista[0]["nome"] == "Lucia"            # repetido não troca um nome já existente

    talvez = (
        "Telefone\n5521986128098\n5521973701889\n5521986128098\n"
        "209083319713962\n"
    )
    r = processar_csv(talvez)
    assert (r.total, len(r.ok), r.duplicados) == (4, 2, 1)
    assert len(r.invalidos) == 1                 # 15 dígitos, não é telefone
    assert r.invalidos[0].motivo == "tamanho fora do padrão"

    # separador ponto e vírgula (Excel em português), com vírgula dentro de aspas
    ponto_virgula = (
        'Nome do Contato;WhatsApp do Contato;Etapa\n'
        '"Silva, João";5521994833600;Fez Contato\n'
        'Ana;(34) 99123-4567;Fez Contato\n'
    )
    r = processar_csv(ponto_virgula)
    assert [(c["nome"], c["numero"]) for c in r.como_lista()] == [
        ("Silva, João", "+5521994833600"), ("Ana", "+5534991234567"),
    ], r.ok
    assert detectar_delimitador(ponto_virgula) == ";"

    # tabulação, coluna única (sem separador) e vírgula continuam funcionando
    assert list(processar_csv("Nome\tTelefone\nLia\t5521986128098\n").ok) == ["+5521986128098"]
    assert list(processar_csv("Telefone\n5521986128098\n").ok) == ["+5521986128098"]
    assert detectar_delimitador("a,b,c\n1;2;3") == ","

    # arquivo salvo pelo Excel (cp1252) com acento
    r = processar_csv("Nome;Telefone\nJosé;5521986128098\n".encode("cp1252"))
    assert r.como_lista()[0]["nome"] == "José"

    # lista só com números, SEM cabeçalho: nenhuma linha pode ser perdida
    r = processar_csv("5521994833600\n5521986128098\n5532988176797\n")
    assert (r.total, len(r.ok)) == (3, 3), r

    # várias colunas: só a do número é usada
    ok = ["+5521994833600", "+5521986128098"]
    r = processar_csv("Número do Pedido,Telefone\n1001,5521994833600\n1002,5521986128098\n")
    assert list(r.ok) == ok
    r = processar_csv("Telefone Fixo,WhatsApp\n(34) 3234-5678,5521994833600\n")
    assert list(r.ok) == ["+5521994833600"] and not r.fixos
    r = processar_csv("ID,Contato\n1,5521994833600\n2,5521986128098\n")  # sem palavra-chave
    assert list(r.ok) == ok
    r = processar_csv("Telefone,Celular\n5521994833600,\n5521986128098,\n")  # Celular vazio
    assert list(r.ok) == ok
    r = processar_csv("Sobrenome,Nome,Tel\nSilva,Ana,5521994833600\n")
    assert r.como_lista()[0] == {"nome": "Ana", "numero": "+5521994833600"}

    # duplicado aproveita o nome da linha repetida; célula vazia é motivo "vazio"
    r = processar_csv("Nome,Telefone\n,5521994833600\nLucia,5521994833600\n")
    assert r.como_lista() == [{"nome": "Lucia", "numero": "+5521994833600"}] and r.duplicados == 1
    r = processar_csv("Nome,Telefone\nMaria,\nJoão,5521994833600\n")
    assert r.invalidos[0].motivo == "vazio" and len(r.ok) == 1

    # arquivo vazio / só cabeçalho não quebram
    assert processar_csv("").total == 0 and processar_csv("Telefone\n").total == 0

    # quebra de linha dentro de aspas e CRLF (Windows)
    r = processar_csv('Nome,Telefone\r\n"Ana\nSilva",5521994833600\r\nBia,5521986128098\r\n')
    assert r.como_lista() == [{"nome": "Ana\nSilva", "numero": "+5521994833600"},
                              {"nome": "Bia", "numero": "+5521986128098"}], r.como_lista()

    # arquivo grande: vários blocos de 1 MB, com acentos que caem nas bordas dos blocos
    N = 80_000
    texto = "Nome,Telefone\n" + "".join(f"José Ângela {i:07d},55219{i:08d}\n" for i in range(N))
    for codificacao in ("utf-8", "cp1252"):
        bruto = texto.encode(codificacao)
        assert len(bruto) > 2 * TAMANHO_BLOCO
        r = processar_csv(bruto)
        assert r.total == N and len(r.ok) == N, (codificacao, r.total)
        assert all(c["nome"] == f"José Ângela {k:07d}" for k, c in enumerate(r.como_lista())), codificacao
    d = previa_csv(texto.encode("utf-8"))
    assert d["aproximado"] and abs(d["total_linhas"] - N) <= 2 and len(d["linhas"]) == 8, d["total_linhas"]

    # preview mostra exatamente como o arquivo será lido
    d = previa_csv("5521994833600\n5521986128098\n")
    assert (d["tem_cabecalho"], d["cabecalho"], d["total_linhas"], d["aproximado"]) == (False, ["Coluna 1"], 2, False)
    d = previa_csv("Nome;WhatsApp\nAna;5521994833600\n")
    assert (d["delimitador"], d["idx_numero"], d["cabecalho"]) == (";", 1, ["Nome", "WhatsApp"])
    assert previa_csv("")["vazio"] is True

    # limite de linhas protege a memória do servidor
    try:
        processar_csv("Telefone\n5521994833600\n5521986128098\n5521973701889\n", max_linhas=2)
        raise AssertionError("deveria ter recusado")
    except LimiteDeLinhas:
        pass

    # ---- planilhas .xlsx (o cliente enviou uma com 5 colunas) ----
    from openpyxl import Workbook

    def xlsx(linhas, abas_extras=False):
        wb = Workbook()
        ws = wb.active
        ws.title = "Contatos"
        for l in linhas:
            ws.append(l)
        if abas_extras:
            wb.create_sheet("Outra").append(["lixo", "lixo"])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    planilha = xlsx([
        ["Nome", "Email", "Telefone", "Cidade", "Etapa"],
        ["Ana", "ana@x.com", 5521994833600, "Rio", "Fez Contato"],         # número como inteiro
        ["Bia", "bia@x.com", "(21) 98612-8098", "Niterói", "Novo"],        # número como texto
        ["Caio", "caio@x.com", 5521986128098.0, "Rio", "Novo"],            # decimal: repete a Bia
        ["Duda", "duda@x.com", "(34) 3234-5678", "Uberlândia", "Novo"],    # fixo
        ["Eli", "eli@x.com", None, "SP", "Novo"],                          # célula vazia
    ], abas_extras=True)
    r = processar_csv(planilha)
    assert [(c["nome"], c["numero"]) for c in r.como_lista()] == [("Ana", "+5521994833600"), ("Bia", "+5521986128098")]
    assert (r.total, r.duplicados, len(r.fixos), len(r.invalidos)) == (5, 1, 1, 1), r
    d = previa_csv(planilha)
    assert (d["formato"], d["aba"], d["idx_numero"], len(d["cabecalho"]), d["aproximado"]) == ("xlsx", "Contatos", 2, 5, False), d

    # xlsx sem cabeçalho e no estilo "já contatados"
    assert len(processar_csv(xlsx([[5521994833600], [5521986128098]])).ok) == 2
    c = processar_csv(xlsx([["Nome do Contato", "WhatsApp do Contato", "Etapa da Jornada"],
                            ["Lucia", 5521994833600, "Fez Contato"]]))
    assert c.como_lista() == [{"nome": "Lucia", "numero": "+5521994833600"}]

    # planilha grande: o total no preview vem do tamanho declarado da aba
    grande = xlsx([["Nome", "Telefone"]] + [[f"N{i}", 5521900000000 + i] for i in range(1500)])
    d = previa_csv(grande)
    assert d["aproximado"] and d["total_linhas"] == 1500, d["total_linhas"]

    # formatos que não dá para ler recebem mensagem clara (e não derrubam o servidor)
    for ruim, trecho in [
        (b"\xd0\xcf\x11\xe0" + b"x" * 50, ".xls antigo"),
        (b"PK\x03\x04 isto nao e uma planilha", "Não foi possível abrir"),
        (b"%PDF-1.4\x00\x01\x02 binario", "não suportado"),
    ]:
        try:
            processar_csv(ruim)
            raise AssertionError("deveria ter recusado")
        except ValueError as e:
            assert trecho in str(e), (trecho, str(e))

    # CSV em UTF-16 ("Texto Unicode" do Excel) também funciona
    utf16 = "Nome\tTelefone\nAna\t5521994833600\n".encode("utf-16")
    assert list(processar_csv(utf16).ok) == ["+5521994833600"]
    print("testes da etapa 2 OK")


if __name__ == "__main__":
    _testes()